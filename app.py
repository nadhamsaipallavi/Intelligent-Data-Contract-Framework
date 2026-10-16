from __future__ import annotations

import io
import json
import re
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
import plotly.express as px
import streamlit as st
from openpyxl.styles import Alignment, Font, PatternFill

st.set_page_config(
    page_title="Intelligent Data Contract Framework",
    page_icon="🏥",
    layout="wide",
    initial_sidebar_state="expanded",
)

PROJECT_TITLE = "Hospital Multi-Format Data Integration, EDA and Preprocessing"

# Dataset files that are shipped with the app (put them in a folder called "data"
# next to this app.py file in your GitHub repo).
DATA_DIR = Path(__file__).parent / "data"
ASSETS_DIR = Path(__file__).parent / "assets"
ARCH_IMAGE = ASSETS_DIR / "architecture.jpg"
SAMPLE_FILES = [
    ("CSV", "hospital_source.csv"),
    ("JSON", "hospital_source.json"),
    ("XML", "hospital_source.xml"),
    ("HTML", "hospital_source.html"),
    ("KML", "hospital_source.kml"),
    ("KMZ", "hospital_source.kmz"),
]

COMMON_COLUMNS = [
    "Hospital_Name", "Address", "City", "ZIP_Code", "Phone", "Website",
    "County_Status", "Longitude", "Latitude", "Location", "Data_Year",
    "Hospital_Type", "Bed_Count", "Quality_Score", "Record_Status",
]
NUMERIC_COLUMNS = ["Longitude", "Latitude", "Data_Year", "Bed_Count", "Quality_Score"]
TEXT_COLUMNS = [
    "Hospital_Name", "Address", "City", "Phone", "Website",
    "County_Status", "Hospital_Type", "Record_Status",
]

FIELD_DESCRIPTIONS = {
    "Hospital_Name": ("Text", "Name of the hospital / facility"),
    "Address": ("Text", "Street address"),
    "City": ("Text", "City / town / municipality"),
    "ZIP_Code": ("Text (5 digits)", "Postal code"),
    "Phone": ("Text (10 digits)", "Contact phone number"),
    "Website": ("Text", "Hospital website / web link"),
    "County_Status": ("Category", "In-County or Out-of-County flag"),
    "Longitude": ("Numeric", "Longitude coordinate (-180 to 180)"),
    "Latitude": ("Numeric", "Latitude coordinate (-90 to 90)"),
    "Location": ("Text", "Raw coordinate string (mainly from KML/KMZ)"),
    "Data_Year": ("Numeric", "Reporting / data year"),
    "Hospital_Type": ("Category", "General Hospital, Emergency Center, Specialty Hospital, etc."),
    "Bed_Count": ("Numeric", "Number of beds (must be >= 0)"),
    "Quality_Score": ("Numeric", "Quality / inspection score (0 to 100)"),
    "Record_Status": ("Text", "Status flag of the record (HTML source)"),
}

# These mappings reproduce the source notebook's field mappings.
SCHEMA_MAPPING = {
    "CSV": {
        "Hospital_Name": "Hospital_Name", "Address": "Address", "City": "City",
        "ZIP_Code": "ZIP_Code", "Phone": "Phone", "Website": "Website",
        "County_Status": "County_Status", "Longitude": "Longitude",
        "Latitude": "Latitude", "Location": "Location", "Data_Year": "Data_Year",
        "Hospital_Type": "Hospital_Type", "Bed_Count": "Bed_Count",
        "Quality_Score": "Inspection_Score", "Record_Status": None,
    },
    "JSON": {
        "Hospital_Name": "hospital_name", "Address": "street_address", "City": "city",
        "ZIP_Code": "postal_code", "Phone": "telephone", "Website": "web",
        "County_Status": "county_status", "Longitude": "lon", "Latitude": "lat",
        "Location": None, "Data_Year": "year", "Hospital_Type": "hospital_type",
        "Bed_Count": "beds", "Quality_Score": "rating", "Record_Status": None,
    },
    "XML": {
        "Hospital_Name": "HospitalName", "Address": "Street", "City": "Town",
        "ZIP_Code": "PostalCode", "Phone": "ContactPhone", "Website": "Website",
        "County_Status": "CountyStatus", "Longitude": "Longitude",
        "Latitude": "Latitude", "Location": None, "Data_Year": "Year",
        "Hospital_Type": "Type", "Bed_Count": "Beds", "Quality_Score": "QualityScore",
        "Record_Status": None,
    },
    "HTML": {
        "Hospital_Name": "Facility Name", "Address": "Street Address",
        "City": "Municipality", "ZIP_Code": "Postal", "Phone": "Contact",
        "Website": "Web Link", "County_Status": "County Flag", "Longitude": "Lon",
        "Latitude": "Lat", "Location": None, "Data_Year": "Reporting Year",
        "Hospital_Type": "Facility Category", "Bed_Count": "Beds Available",
        "Quality_Score": None, "Record_Status": "Record Status",
    },
    "KML": {
        "Hospital_Name": "name", "Address": None, "City": None, "ZIP_Code": None,
        "Phone": None, "Website": None, "County_Status": None, "Longitude": None,
        "Latitude": None, "Location": "coordinates", "Data_Year": None,
        "Hospital_Type": None, "Bed_Count": None, "Quality_Score": None,
        "Record_Status": None,
    },
    "KMZ": {
        "Hospital_Name": "name", "Address": None, "City": None, "ZIP_Code": None,
        "Phone": None, "Website": None, "County_Status": None, "Longitude": None,
        "Latitude": None, "Location": "coordinates", "Data_Year": None,
        "Hospital_Type": None, "Bed_Count": None, "Quality_Score": None,
        "Record_Status": None,
    },
}

ALIASES = {
    "Hospital_Name": ["hospital name", "facility name", "name", "hospital_name", "hospitalname"],
    "Address": ["address", "street", "street address", "street_address"],
    "City": ["city", "town", "municipality"],
    "ZIP_Code": ["zip", "zip code", "postal", "postal code", "postal_code", "zipcode"],
    "Phone": ["phone", "telephone", "contact", "contactphone", "contact phone"],
    "Website": ["website", "web", "web link", "url"],
    "County_Status": ["county status", "county_status", "county flag"],
    "Longitude": ["longitude", "lon", "lng"],
    "Latitude": ["latitude", "lat"],
    "Location": ["location", "coordinates"],
    "Data_Year": ["data year", "year", "reporting year"],
    "Hospital_Type": ["hospital type", "hospital_type", "type", "facility category"],
    "Bed_Count": ["bed count", "beds", "bed_count", "beds available"],
    "Quality_Score": ["quality score", "qualityscore", "rating", "inspection score", "inspection_score"],
    "Record_Status": ["record status", "record_status"],
}


# ---------------------------------------------------------------------
# Core processing functions (unchanged logic)
# ---------------------------------------------------------------------
def norm_name(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value).strip().lower())


def detect_format(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    mapping = {".csv": "CSV", ".json": "JSON", ".xml": "XML", ".html": "HTML",
               ".htm": "HTML", ".kml": "KML", ".kmz": "KMZ"}
    if ext not in mapping:
        raise ValueError(f"Unsupported file extension: {ext or '(none)'}.")
    return mapping[ext]


def parse_kml_bytes(data: bytes) -> pd.DataFrame:
    root = ET.fromstring(data)
    rows = []
    for placemark in root.iter():
        if placemark.tag.split("}")[-1] != "Placemark":
            continue
        row = {}
        for child in placemark.iter():
            tag = child.tag.split("}")[-1]
            if tag == "name" and child.text and "name" not in row:
                row["name"] = child.text.strip()
            elif tag == "description" and child.text:
                row["description"] = child.text.strip()
            elif tag == "coordinates" and child.text:
                row["coordinates"] = child.text.strip()
        row.setdefault("coordinates", np.nan)
        rows.append(row)
    return pd.DataFrame(rows)


def parse_upload(uploaded_file, source_format: str) -> pd.DataFrame:
    data = uploaded_file.getvalue()
    if source_format == "CSV":
        return pd.read_csv(io.BytesIO(data))
    if source_format == "JSON":
        obj = json.loads(data.decode("utf-8-sig"))
        if isinstance(obj, dict):
            candidate = next((v for v in obj.values() if isinstance(v, list)), None)
            obj = candidate if candidate is not None else [obj]
        return pd.json_normalize(obj)
    if source_format == "XML":
        root = ET.fromstring(data)
        records = []
        for element in root.iter():
            if element.tag.split("}")[-1].lower() in {"hospital", "record", "facility"}:
                record = {}
                for child in list(element):
                    record[child.tag.split("}")[-1]] = child.text
                if record:
                    records.append(record)
        if not records:
            for element in list(root):
                record = {child.tag.split("}")[-1]: child.text for child in list(element)}
                if record:
                    records.append(record)
        return pd.DataFrame(records)
    if source_format == "HTML":
        tables = pd.read_html(io.BytesIO(data))
        if not tables:
            raise ValueError("No HTML table was found in this file.")
        return tables[0].copy()
    if source_format == "KML":
        return parse_kml_bytes(data)
    if source_format == "KMZ":
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            kml_names = [n for n in archive.namelist()
                         if n.lower().endswith(".kml") and not n.startswith("__MACOSX/")]
            if not kml_names:
                raise ValueError("No .kml file was found inside this KMZ archive.")
            return parse_kml_bytes(archive.read(kml_names[0]))
    raise ValueError(f"Unsupported format: {source_format}")


class LocalFile:
    """Wraps a file on disk so it behaves like a Streamlit uploaded file."""

    def __init__(self, path: Path):
        self.path = Path(path)
        self.name = self.path.name

    def getvalue(self) -> bytes:
        return self.path.read_bytes()


def sample_files() -> list[LocalFile]:
    return [LocalFile(DATA_DIR / name) for _, name in SAMPLE_FILES if (DATA_DIR / name).exists()]


@st.cache_data(show_spinner=False)
def load_sample_raw(filename: str) -> pd.DataFrame:
    f = LocalFile(DATA_DIR / filename)
    return parse_upload(f, detect_format(filename))


def standardize_schema(df: pd.DataFrame, source_format: str) -> tuple[pd.DataFrame, pd.DataFrame]:
    mapping = SCHEMA_MAPPING[source_format]
    lookup = {norm_name(c): c for c in df.columns}
    standardized = pd.DataFrame(index=df.index)
    mapping_rows = []
    for target in COMMON_COLUMNS:
        source = mapping.get(target)
        chosen = source if source is not None and source in df.columns else None
        if chosen is None:
            target_norm = norm_name(target)
            if target_norm in lookup:
                chosen = lookup[target_norm]
            else:
                for alias in ALIASES.get(target, []):
                    if norm_name(alias) in lookup:
                        chosen = lookup[norm_name(alias)]
                        break
        standardized[target] = df[chosen] if chosen is not None else np.nan
        mapping_rows.append({
            "Source_Format": source_format, "Common_Field": target,
            "Source_Field": chosen if chosen is not None else "Not available",
            "Mapping_Status": "Mapped" if chosen is not None else "Missing in source",
        })
    standardized["Source_Format"] = source_format
    standardized["Source_Record_ID"] = range(1, len(standardized) + 1)
    return standardized, pd.DataFrame(mapping_rows)


def make_missing_report(df: pd.DataFrame) -> pd.DataFrame:
    report = pd.DataFrame({
        "Column": df.columns,
        "Missing_Count": [int(df[c].isna().sum()) for c in df.columns],
    })
    report["Missing_Percentage"] = (
        report["Missing_Count"] / max(len(df), 1) * 100
    ).round(2)
    return report.sort_values("Missing_Count", ascending=False)


def preprocess_data(raw: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, dict]:
    clean = raw.copy()
    before_rows = len(clean)
    before_missing = int(clean.isna().sum().sum())
    before_duplicates = int(clean.duplicated().sum())
    actions = []

    for col in TEXT_COLUMNS:
        clean[col] = clean[col].astype("string").str.strip()
    actions.append({"Step": "Text standardization", "Action": "Trimmed whitespace from text fields"})

    county_mapping = {
        "in county": "In-County", "in-county": "In-County",
        "out county": "Out-of-County", "out-of-county": "Out-of-County",
    }
    clean["County_Status"] = clean["County_Status"].apply(
        lambda x: county_mapping.get(str(x).strip().lower(), x) if pd.notna(x) else x
    )
    actions.append({"Step": "County category standardization", "Action": "Normalized county status labels"})

    type_mapping = {
        "general": "General Hospital", "emergency": "Emergency Center",
        "specialty": "Specialty Hospital",
    }
    clean["Hospital_Type"] = clean["Hospital_Type"].apply(
        lambda x: type_mapping.get(str(x).strip().lower(), x) if pd.notna(x) else x
    )
    actions.append({"Step": "Hospital type standardization", "Action": "Normalized common hospital type labels"})

    for col in NUMERIC_COLUMNS:
        clean[col] = pd.to_numeric(clean[col], errors="coerce")
    actions.append({"Step": "Numeric conversion", "Action": "Converted numeric fields; invalid text becomes missing"})

    invalid_beds = int((clean["Bed_Count"] < 0).sum())
    clean.loc[clean["Bed_Count"] < 0, "Bed_Count"] = np.nan
    invalid_quality = int(((clean["Quality_Score"] < 0) | (clean["Quality_Score"] > 100)).sum())
    clean.loc[(clean["Quality_Score"] < 0) | (clean["Quality_Score"] > 100), "Quality_Score"] = np.nan

    invalid_lon = (clean["Longitude"] < -180) | (clean["Longitude"] > 180)
    invalid_lat = (clean["Latitude"] < -90) | (clean["Latitude"] > 90)
    coordinate_errors = int(invalid_lon.sum() + invalid_lat.sum())
    clean.loc[invalid_lon, "Longitude"] = np.nan
    clean.loc[invalid_lat, "Latitude"] = np.nan

    clean["ZIP_Code"] = clean["ZIP_Code"].astype("string").str.strip().str.replace(r"\.0$", "", regex=True)
    zip_valid = clean["ZIP_Code"].str.match(r"^\d{5}$", na=False)
    invalid_zip = int((~zip_valid & clean["ZIP_Code"].notna()).sum())
    clean.loc[~zip_valid, "ZIP_Code"] = pd.NA

    def clean_phone(value):
        if pd.isna(value):
            return pd.NA
        digits = re.sub(r"\D", "", str(value))
        return digits if len(digits) == 10 else pd.NA

    clean["Phone"] = clean["Phone"].apply(clean_phone)
    clean["Website"] = clean["Website"].astype("string").str.strip().replace("", pd.NA)
    clean["Hospital_Name"] = clean["Hospital_Name"].astype("string").str.strip().str.replace(
        "National Institue", "National Institute", regex=False
    )

    rows_before_dedup = len(clean)
    clean = clean.drop_duplicates().reset_index(drop=True)
    duplicates_removed = rows_before_dedup - len(clean)

    imputation_rows = []
    for col in ["Bed_Count", "Quality_Score"]:
        missing_before = int(clean[col].isna().sum())
        median = clean[col].median()
        if pd.notna(median):
            clean[col] = clean[col].fillna(median)
        imputation_rows.append({
            "Column": col, "Missing_Before": missing_before,
            "Imputation_Method": "Median",
            "Imputation_Value": median if pd.notna(median) else "No median available",
            "Missing_After": int(clean[col].isna().sum()),
        })
    imputation_report = pd.DataFrame(imputation_rows)

    validation = []
    checks = [
        ("Negative Bed Count", int((clean["Bed_Count"] < 0).sum())),
        ("Quality Score 0-100", int(((clean["Quality_Score"] < 0) | (clean["Quality_Score"] > 100)).sum())),
        ("Latitude Range", int(((clean["Latitude"] < -90) | (clean["Latitude"] > 90)).sum())),
        ("Longitude Range", int(((clean["Longitude"] < -180) | (clean["Longitude"] > 180)).sum())),
        ("Duplicate Rows", int(clean.duplicated().sum())),
        ("Invalid ZIP format", int((clean["ZIP_Code"].notna() & ~clean["ZIP_Code"].str.match(r"^\d{5}$", na=False)).sum())),
        ("Invalid phone format", int((clean["Phone"].notna() & ~clean["Phone"].str.match(r"^\d{10}$", na=False)).sum())),
    ]
    for name, count in checks:
        validation.append({"Validation_Check": name, "Invalid_Count": count,
                           "Status": "PASS" if count == 0 else "FAIL"})
    validation_report = pd.DataFrame(validation)

    actions.extend([
        {"Step": "Invalid values", "Action": f"Invalid beds: {invalid_beds}; invalid quality scores: {invalid_quality}; invalid coordinate values: {coordinate_errors} set to missing"},
        {"Step": "ZIP validation", "Action": f"{invalid_zip} invalid ZIP values set to missing"},
        {"Step": "Phone standardization", "Action": "Kept only 10-digit phone numbers; others set to missing"},
        {"Step": "Hospital name standardization", "Action": "Applied notebook spelling correction for 'National Institue'"},
        {"Step": "Duplicate removal", "Action": f"Removed {duplicates_removed} exact duplicate rows"},
        {"Step": "Median imputation", "Action": "Filled missing Bed_Count and Quality_Score with column medians when available"},
    ])
    stats = {
        "rows_before": before_rows, "rows_after": len(clean),
        "columns": len(clean.columns), "missing_before": before_missing,
        "missing_after": int(clean.isna().sum().sum()),
        "duplicates_before": before_duplicates, "duplicates_removed": duplicates_removed,
        "invalid_beds": invalid_beds, "invalid_quality": invalid_quality,
        "invalid_coordinates": coordinate_errors, "invalid_zip": invalid_zip,
    }
    return clean, validation_report, imputation_report, {"stats": stats, "actions": pd.DataFrame(actions)}


def create_excel_report(raw, clean, source_summary, schema_report, missing_report,
                        validation_report, imputation_report, actions, before_after,
                        dtype_report, outlier_report) -> bytes:
    final_display = clean.copy().fillna("NA")
    out = io.BytesIO()
    sheets = {
        "Raw_Combined_Data": raw, "Final_Clean_Data": final_display,
        "Source_Summary": source_summary, "Schema_Mapping": schema_report,
        "Data_Types": dtype_report, "Missing_Values": missing_report,
        "Outlier_Analysis": outlier_report, "Before_After": before_after,
        "Preprocessing_Log": actions, "Validation_Report": validation_report,
        "Imputation_Report": imputation_report,
    }
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        for sheet, df in sheets.items():
            df.to_excel(writer, sheet_name=sheet[:31], index=False)
        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="D9EAF7")
                cell.alignment = Alignment(horizontal="center", vertical="center")
            for column_cells in ws.columns:
                max_len = max((len(str(c.value)) for c in column_cells if c.value is not None), default=10)
                ws.column_dimensions[column_cells[0].column_letter].width = min(max_len + 2, 35)
            for row in ws.iter_rows():
                for cell in row:
                    if str(cell.value).strip().upper() == "NA":
                        cell.fill = PatternFill("solid", fgColor="FFF2CC")
    return out.getvalue()


def process_files(uploads) -> bool:
    """Parse, standardize, preprocess and store results in session state."""
    parsed, standardized, mapping_frames, errors = [], [], [], []
    progress = st.progress(0, text="Reading files…")
    for i, file in enumerate(uploads):
        try:
            fmt = detect_format(file.name)
            df = parse_upload(file, fmt)
            if df.empty:
                raise ValueError("The file was read, but no records were found.")
            std, mapping = standardize_schema(df, fmt)
            parsed.append({"filename": file.name, "format": fmt, "raw": df, "standardized": std})
            standardized.append(std)
            mapping_frames.append(mapping.assign(Source_File=file.name))
        except Exception as exc:
            errors.append({"File": file.name, "Error": str(exc)})
        progress.progress((i + 1) / len(uploads), text=f"Processed {i + 1} of {len(uploads)} file(s)")
    progress.empty()
    if errors:
        st.warning("Some files could not be processed. Review the details below; successfully read files are still available.")
        st.dataframe(pd.DataFrame(errors), use_container_width=True, hide_index=True)
    if not standardized:
        st.error("No files were processed successfully. Check file formats and structure.")
        return False

    raw = pd.concat(standardized, ignore_index=True)
    mapping_report = pd.concat(mapping_frames, ignore_index=True)
    source_summary = pd.DataFrame([
        {"File": p["filename"], "Source_Format": p["format"],
         "Input_Rows": len(p["raw"]), "Input_Columns": len(p["raw"].columns),
         "Mapped_Fields": int((mapping_report.loc[mapping_report["Source_File"] == p["filename"], "Mapping_Status"] == "Mapped").sum())}
        for p in parsed
    ])
    clean, validation_report, imputation_report, metadata = preprocess_data(raw)
    stats = metadata["stats"]
    missing_report = make_missing_report(raw)
    dtype_report = pd.DataFrame({"Column": raw.columns, "Data_Type": raw.dtypes.astype(str).values})
    outlier_rows = []
    for col in ["Bed_Count", "Quality_Score"]:
        vals = pd.to_numeric(raw[col], errors="coerce").dropna()
        if len(vals):
            q1, q3 = vals.quantile(.25), vals.quantile(.75)
            iqr = q3 - q1
            low, high = q1 - 1.5 * iqr, q3 + 1.5 * iqr
            outlier_rows.append({"Variable": col, "Outlier_Count": int(((vals < low) | (vals > high)).sum()),
                                 "Lower_Bound": low, "Upper_Bound": high})
        else:
            outlier_rows.append({"Variable": col, "Outlier_Count": 0, "Lower_Bound": np.nan, "Upper_Bound": np.nan})
    outlier_report = pd.DataFrame(outlier_rows)
    before_after = pd.DataFrame([
        {"Metric": "Rows", "Before": len(raw), "After": len(clean)},
        {"Metric": "Missing cells", "Before": stats["missing_before"], "After": stats["missing_after"]},
        {"Metric": "Exact duplicate rows", "Before": stats["duplicates_before"], "After": int(clean.duplicated().sum())},
    ])
    st.session_state["hospital_result"] = {
        "parsed": parsed, "raw": raw, "clean": clean, "mapping": mapping_report,
        "source_summary": source_summary, "missing": missing_report, "dtypes": dtype_report,
        "validation": validation_report, "imputation": imputation_report,
        "actions": metadata["actions"], "before_after": before_after,
        "outliers": outlier_report, "errors": errors,
    }
    return True


# ---------------------------------------------------------------------
# Styling
# ---------------------------------------------------------------------
st.markdown("""
<style>
.block-container {padding-top: 1.5rem; padding-bottom: 3rem; max-width: 1500px;}
.hero {padding: 1.3rem 1.5rem; border-radius: 18px; background: linear-gradient(120deg,#102a43,#176b87); color: white; margin-bottom: 1.2rem;}
.hero h1 {font-size: 2.1rem; margin: 0 0 .4rem 0;}
.hero p {margin: 0; opacity: .9;}
.card {border: 1px solid rgba(128,128,128,.25); border-radius: 14px; padding: 1rem 1.2rem; margin-bottom: 1rem;}
.card h4 {margin-top: 0;}
div[data-testid="stMetric"] {border: 1px solid rgba(128,128,128,.22); padding: 14px; border-radius: 14px;}

/* ---------- About page cards ---------- */
.feature-card {border: 1px solid rgba(128,128,128,.25); border-radius: 14px; padding: 1rem 1.1rem; height: 100%;
               border-top: 4px solid #176b87; margin-bottom: .8rem;}
.feature-card .icon {font-size: 1.6rem; margin-bottom: .3rem;}
.feature-card h5 {margin: 0 0 .35rem 0; font-size: 1.02rem;}
.feature-card p {margin: 0; font-size: .92rem; opacity: .85; line-height: 1.5;}
.stage {display: flex; gap: .9rem; align-items: flex-start; padding: .8rem 1rem; margin-bottom: .6rem;
        border: 1px solid rgba(128,128,128,.22); border-radius: 12px;}
.stage .num {min-width: 34px; height: 34px; border-radius: 50%; background: linear-gradient(135deg,#102a43,#176b87);
             color: #fff; display: flex; align-items: center; justify-content: center; font-weight: 700;}
.stage .body b {display: block; margin-bottom: .15rem;}
.stage .body span {font-size: .92rem; opacity: .85; line-height: 1.5;}
.badge {display: inline-block; padding: .18rem .65rem; margin: .15rem .2rem .15rem 0; border-radius: 999px;
        background: rgba(23,107,135,.14); border: 1px solid rgba(23,107,135,.35); font-size: .85rem;}

/* ---------- Architecture page ---------- */
.arch-flow {display: flex; flex-direction: column; align-items: stretch; gap: 0; margin: 1rem 0 1.4rem 0;}
.arch-layer {border: 1px solid rgba(128,128,128,.28); border-left: 6px solid #176b87; border-radius: 12px;
             padding: .85rem 1.1rem;}
.arch-layer .t {font-weight: 700; font-size: 1.02rem; margin-bottom: .2rem;}
.arch-layer .d {font-size: .9rem; opacity: .85; line-height: 1.5;}
.arch-layer .f {font-size: .78rem; opacity: .65; margin-top: .3rem; font-family: monospace;}
.arch-arrow {text-align: center; font-size: 1.3rem; line-height: 1.5rem; color: #176b87;}
.arch-chips {margin-top: .45rem;}

/* ---------- Sidebar menu ---------- */
section[data-testid="stSidebar"] {background: linear-gradient(180deg,#0b1f33 0%,#102a43 55%,#143d57 100%);}
section[data-testid="stSidebar"] * {color: #e8f1f7;}
section[data-testid="stSidebar"] hr {border-color: rgba(255,255,255,.14);}
.side-brand {display: flex; align-items: center; gap: .7rem; padding: .4rem .2rem 1rem .2rem;}
.side-brand .logo {width: 46px; height: 46px; border-radius: 14px; font-size: 1.6rem;
                   background: linear-gradient(135deg,#1fa2c7,#176b87); display: flex;
                   align-items: center; justify-content: center; box-shadow: 0 4px 14px rgba(0,0,0,.35);}
.side-brand .title {font-weight: 700; font-size: 1.05rem; line-height: 1.15;}
.side-brand .sub {font-size: .74rem; opacity: .7; letter-spacing: .04em;}
.side-label {font-size: .72rem; letter-spacing: .12em; text-transform: uppercase; opacity: .6; margin: .2rem 0 .5rem .3rem;}

section[data-testid="stSidebar"] div[role="radiogroup"] {gap: .35rem;}
section[data-testid="stSidebar"] div[role="radiogroup"] > label {
    width: 100%; padding: .65rem .85rem; border-radius: 12px; cursor: pointer;
    border: 1px solid transparent; transition: all .18s ease;
    background: rgba(255,255,255,.04);
}
section[data-testid="stSidebar"] div[role="radiogroup"] > label:hover {
    background: rgba(255,255,255,.12); transform: translateX(3px);
}
/* hide the round radio dot */
section[data-testid="stSidebar"] div[role="radiogroup"] > label > div:first-child {display: none;}
section[data-testid="stSidebar"] div[role="radiogroup"] > label p {font-size: .98rem; font-weight: 500; margin: 0;}
/* active page */
section[data-testid="stSidebar"] div[role="radiogroup"] > label:has(input:checked) {
    background: linear-gradient(90deg,#1fa2c7,#176b87);
    border-color: rgba(255,255,255,.35);
    box-shadow: 0 6px 18px rgba(0,0,0,.35);
}
section[data-testid="stSidebar"] div[role="radiogroup"] > label:has(input:checked) p {font-weight: 700; color: #fff;}

.status-pill {display: flex; align-items: center; gap: .55rem; padding: .6rem .8rem; border-radius: 12px;
              background: rgba(255,255,255,.06); border: 1px solid rgba(255,255,255,.12); font-size: .85rem;}
.dot {width: 10px; height: 10px; border-radius: 50%; display: inline-block;}
.dot.on {background: #3ddc97; box-shadow: 0 0 8px #3ddc97;}
.dot.off {background: #f5b942; box-shadow: 0 0 8px #f5b942;}
.side-foot {font-size: .75rem; opacity: .6; margin-top: .8rem; line-height: 1.5;}
</style>
<div class="hero">
<h1>🏥 Intelligent Data Contract Framework</h1>
<p>Integrate heterogeneous hospital data, explore data quality, preprocess records, validate, and export.</p>
</div>
""", unsafe_allow_html=True)


def require_result():
    result = st.session_state.get("hospital_result")
    if not result:
        st.info("No results yet. Open **▶️ Implementation** in the sidebar, run the bundled dataset files "
                "(or upload your own) and the results will appear here.")
    return result


# ---------------------------------------------------------------------
# 1. ABOUT PROJECT (shorter version)
# ---------------------------------------------------------------------
def page_about_project():
    st.header("ℹ️ About the Project")

    st.markdown(f"""
<div class="card">
<h4>{PROJECT_TITLE}</h4>
Hospital data is published by different agencies, each with its own file format, column names and quality.
This application <b>integrates</b> hospital records from <b>six formats</b> (CSV, JSON, XML, HTML, KML and KMZ)
into one common <b>15-field schema</b>, explores data quality with <b>EDA</b>, cleans it with a documented
<b>preprocessing</b> pipeline, <b>validates</b> the result and <b>exports</b> analysis-ready data with a full audit report.
</div>
""", unsafe_allow_html=True)

    st.subheader("🧩 Problem and motivation")
    p1, p2 = st.columns(2)
    with p1:
        st.markdown("**The problem**")
        st.markdown(
            "- The same attribute has different names in different files\n"
            "- Values are inconsistent: bad ZIP codes, phone formats, category labels\n"
            "- Impossible values such as negative beds or invalid coordinates\n"
            "- Some sources do not contain certain fields at all"
        )
    with p2:
        st.markdown("**Why it matters**")
        st.markdown(
            "- Analysts spend most of their time cleaning data\n"
            "- Capacity and quality decisions need correct numbers\n"
            "- A repeatable pipeline is more reliable than manual fixes\n"
            "- Every step is transparent and reviewable"
        )

    st.subheader("🎯 Objectives")
    objectives = [
        ("🔗", "Integrate", "Combine six file formats into one standard schema."),
        ("🔍", "Profile", "Expose missing values, types, duplicates and outliers."),
        ("🧹", "Clean", "Standardize and fix data using documented rules."),
        ("✅", "Validate", "Run automated PASS / FAIL checks after cleaning."),
        ("📦", "Deliver", "Export cleaned data and a multi-sheet Excel report."),
        ("🔁", "Reproduce", "Re-run the whole process on new files in one click."),
    ]
    for row_start in range(0, len(objectives), 3):
        cols = st.columns(3)
        for col, (icon, title, text) in zip(cols, objectives[row_start:row_start + 3]):
            with col:
                st.markdown(
                    f'<div class="feature-card"><div class="icon">{icon}</div><h5>{title}</h5><p>{text}</p></div>',
                    unsafe_allow_html=True,
                )

    st.subheader("✨ Key features")
    f1, f2, f3 = st.columns(3)
    features = [
        (f1, "📂 Multi-format ingestion", "Use the bundled dataset or upload your own files."),
        (f2, "🧭 Transparent mapping", "Shows which source column fed each common field."),
        (f3, "📊 Interactive EDA", "Plotly charts and tables for data-quality review."),
    ]
    for col, title, text in features:
        with col:
            st.markdown(f'<div class="feature-card"><h5>{title}</h5><p>{text}</p></div>',
                        unsafe_allow_html=True)

    st.subheader("🧰 Technologies used")
    st.markdown("".join(
        f'<span class="badge">{t}</span>'
        for t in ["Python", "Streamlit", "Pandas", "NumPy", "Plotly", "OpenPyXL", "lxml / html5lib"]
    ), unsafe_allow_html=True)

    st.markdown("**Supported input formats**")
    st.markdown("".join(f'<span class="badge">{fmt}</span>' for fmt, _ in SAMPLE_FILES), unsafe_allow_html=True)


# ---------------------------------------------------------------------
# 2. ARCHITECTURE (now directly below About Project)
# ---------------------------------------------------------------------
def page_architecture():
    st.header("🏗️ Architecture")
    st.write("The application is a layered pipeline. Each layer takes the output of the previous one, "
             "so every step stays separate, testable and traceable.")

    st.subheader("System architecture diagram")
    if ARCH_IMAGE.exists():
        st.image(str(ARCH_IMAGE), use_container_width=True,
                 caption="Left: detailed processing flow. Right: high-level layered architecture.")
        with open(ARCH_IMAGE, "rb") as f:
            st.download_button("⬇️ Download architecture diagram", f.read(),
                               file_name="architecture.jpg", mime="image/jpeg")
    else:
        st.warning("Architecture image not found. Add it to your repository as `assets/architecture.jpg`.")

    st.subheader("Design principles")
    st.dataframe(pd.DataFrame([
        {"Principle": "Traceability", "Meaning": "Every record keeps Source_Format and Source_Record_ID"},
        {"Principle": "No invented data", "Meaning": "Missing descriptive fields stay missing; only numeric gaps use median imputation"},
        {"Principle": "Fail-safe parsing", "Meaning": "A bad file is reported, and the remaining files are still processed"},
        {"Principle": "Auditability", "Meaning": "Logs, before/after counts and validation results are exported"},
    ]), use_container_width=True, hide_index=True)


# ---------------------------------------------------------------------
# 3. ABOUT DATASET
# ---------------------------------------------------------------------
def page_about_dataset():
    st.header("📚 About the Dataset")
    st.write("The project uses hospital records supplied as **six data files**, one per format. "
             "Each file describes the same kind of facilities but uses its own structure and column names. "
             "Every source is converted into one **common 15-field schema**.")

    st.subheader("Dataset files")
    available = {name: (DATA_DIR / name).exists() for _, name in SAMPLE_FILES}
    if not all(available.values()):
        missing = [n for n, ok in available.items() if not ok]
        st.warning("Some dataset files were not found in the `data/` folder: " + ", ".join(missing))

    rows = []
    for fmt, name in SAMPLE_FILES:
        path = DATA_DIR / name
        if not path.exists():
            continue
        try:
            df = load_sample_raw(name)
            rows.append({"Format": fmt, "File": name, "Size (KB)": round(path.stat().st_size / 1024, 1),
                         "Records": len(df), "Fields in source": len(df.columns)})
        except Exception as exc:
            rows.append({"Format": fmt, "File": name, "Size (KB)": round(path.stat().st_size / 1024, 1),
                         "Records": "error", "Fields in source": str(exc)})
    if rows:
        summary = pd.DataFrame(rows)
        st.dataframe(summary, use_container_width=True, hide_index=True)
        numeric_rows = summary[pd.to_numeric(summary["Records"], errors="coerce").notna()]
        if not numeric_rows.empty:
            st.caption(f"Total records across all files: {int(pd.to_numeric(numeric_rows['Records']).sum()):,}")

    st.subheader("Preview each file")
    tabs = st.tabs([fmt for fmt, _ in SAMPLE_FILES])
    notes = {
        "CSV": "Tabular rows and columns. Quality score comes from `Inspection_Score`.",
        "JSON": "List of records. Short keys such as `lat`, `lon`, `beds`, `rating`.",
        "XML": "Repeated `<Hospital>` elements; child elements become fields.",
        "HTML": "First HTML table in the file. Has `Record Status` but no quality score.",
        "KML": "Map placemarks. Only the name and coordinates are used.",
        "KMZ": "Zipped KML file. Same as KML after unzipping.",
    }
    for tab, (fmt, name) in zip(tabs, SAMPLE_FILES):
        with tab:
            st.caption(notes[fmt])
            if (DATA_DIR / name).exists():
                try:
                    st.dataframe(load_sample_raw(name).head(50), use_container_width=True, hide_index=True)
                    st.caption("Showing the first 50 records exactly as they appear in the source file.")
                except Exception as exc:
                    st.error(f"Could not read {name}: {exc}")
            else:
                st.warning(f"`data/{name}` not found.")

    st.subheader("Common schema (15 fields)")
    st.dataframe(pd.DataFrame([
        {"Field": f, "Type": FIELD_DESCRIPTIONS[f][0], "Description": FIELD_DESCRIPTIONS[f][1]}
        for f in COMMON_COLUMNS
    ]), use_container_width=True, hide_index=True)

    st.subheader("Source field mapping")
    st.caption("Which source column feeds each common field. “—” means the source does not provide that field.")
    mapping_table = pd.DataFrame({
        fmt: [SCHEMA_MAPPING[fmt].get(f) or "—" for f in COMMON_COLUMNS]
        for fmt in SCHEMA_MAPPING
    }, index=COMMON_COLUMNS)
    mapping_table.index.name = "Common field"
    st.dataframe(mapping_table, use_container_width=True)

    st.info("If a header does not match exactly, the app also tries normalized names and common aliases "
            "(e.g. “postal code”, “zipcode”). Fields that do not exist in a source are left missing, never invented.")


# ---------------------------------------------------------------------
# 4. PREPROCESSING
# ---------------------------------------------------------------------
def page_preprocessing():
    st.header("🧹 Preprocessing")
    tab_rules, tab_results = st.tabs(["📖 Rules & method", "📊 Results (after implementation)"])

    with tab_rules:
        st.write("Cleaning follows the rules of the source notebook. Descriptive fields such as phone numbers and "
                 "addresses are never fabricated.")
        st.dataframe(pd.DataFrame([
            {"#": 1, "Step": "Trim text", "Rule": "Remove leading/trailing spaces from text fields"},
            {"#": 2, "Step": "County status", "Rule": "“in county” → In-County; “out county” → Out-of-County"},
            {"#": 3, "Step": "Hospital type", "Rule": "general → General Hospital; emergency → Emergency Center; specialty → Specialty Hospital"},
            {"#": 4, "Step": "Numeric conversion", "Rule": "Longitude, Latitude, Data_Year, Bed_Count, Quality_Score converted to numbers; bad text → missing"},
            {"#": 5, "Step": "Invalid beds", "Rule": "Negative Bed_Count → missing"},
            {"#": 6, "Step": "Invalid quality score", "Rule": "Quality_Score outside 0–100 → missing"},
            {"#": 7, "Step": "Coordinate check", "Rule": "Latitude outside −90..90 or Longitude outside −180..180 → missing"},
            {"#": 8, "Step": "ZIP validation", "Rule": "Keep only 5-digit ZIP codes (removes “.0” suffix first)"},
            {"#": 9, "Step": "Phone standardization", "Rule": "Keep digits only; valid only if exactly 10 digits"},
            {"#": 10, "Step": "Name correction", "Rule": "“National Institue” → “National Institute”"},
            {"#": 11, "Step": "Duplicate removal", "Rule": "Remove exact duplicate rows"},
            {"#": 12, "Step": "Median imputation", "Rule": "Fill missing Bed_Count and Quality_Score with the column median"},
        ]), use_container_width=True, hide_index=True)

        st.subheader("Validation checks after cleaning")
        st.dataframe(pd.DataFrame([
            {"Check": "Negative Bed Count", "Passes when": "No negative bed counts"},
            {"Check": "Quality Score 0-100", "Passes when": "All scores within 0–100"},
            {"Check": "Latitude Range", "Passes when": "All latitudes within −90..90"},
            {"Check": "Longitude Range", "Passes when": "All longitudes within −180..180"},
            {"Check": "Duplicate Rows", "Passes when": "No duplicate rows remain"},
            {"Check": "Invalid ZIP format", "Passes when": "All non-missing ZIPs are 5 digits"},
            {"Check": "Invalid phone format", "Passes when": "All non-missing phones are 10 digits"},
        ]), use_container_width=True, hide_index=True)

    with tab_results:
        result = require_result()
        if not result:
            return
        clean = result["clean"]
        st.subheader("Before vs after")
        st.dataframe(result["before_after"], use_container_width=True, hide_index=True)
        st.subheader("Cleaned data")
        st.dataframe(clean.fillna("NA"), use_container_width=True, hide_index=True, height=430)
        st.subheader("Preprocessing log")
        st.dataframe(result["actions"], use_container_width=True, hide_index=True)
        st.subheader("Median imputation report")
        st.dataframe(result["imputation"], use_container_width=True, hide_index=True)
        st.download_button("Download cleaned CSV", clean.fillna("NA").to_csv(index=False).encode("utf-8-sig"),
                           "hospital_final_clean_data.csv", "text/csv")


# ---------------------------------------------------------------------
# 5. EDA
# ---------------------------------------------------------------------
MISSING_TOKENS = ["", " ", "NA", "N/A", "NULL", "null", "None", "NONE",
                  "UNKNOWN", "Unknown", "unknown"]


def build_eda_results(raw: pd.DataFrame) -> dict:
    """Reproduces the EDA section of the Colab notebook on the combined, pre-cleaning data."""
    eda = raw.copy()
    for col in eda.columns:                       # Cell 20 - standardize missing tokens
        if not pd.api.types.is_numeric_dtype(eda[col]):
            eda[col] = eda[col].replace(MISSING_TOKENS, np.nan)
    eda.insert(0, "Combined_Record_ID", range(1, len(eda) + 1))

    res = {"rows": eda.shape[0], "columns": eda.shape[1]}
    res["dtypes"] = pd.DataFrame({"Column": eda.columns, "Data_Type": eda.dtypes.astype(str).values})
    missing = pd.DataFrame({"Column": eda.columns, "Missing_Count": eda.isnull().sum().values})
    missing["Missing_Percentage"] = (missing["Missing_Count"] / max(len(eda), 1) * 100).round(2)
    res["missing"] = missing.sort_values("Missing_Count", ascending=False)
    res["missing_by_source"] = (
        eda.isnull().sum(axis=1).groupby(eda["Source_Format"]).sum()
        .reset_index(name="Total_Missing_Cells")
    )
    res["exact_duplicates"] = int(eda.duplicated().sum())
    res["potential_duplicates"] = int(eda.duplicated(subset=["Hospital_Name", "Address", "City"]).sum())
    res["unique"] = pd.DataFrame({
        "Column": eda.columns,
        "Unique_Values": [int(eda[c].nunique(dropna=True)) for c in eda.columns],
    })
    res["categorical"] = {
        c: eda[c].fillna("NaN (missing)").astype(str).value_counts()
            .rename_axis(c).reset_index(name="Count")
        for c in ["County_Status", "Hospital_Type", "Record_Status", "Source_Format"]
    }

    numeric_cols = ["Longitude", "Latitude", "Data_Year", "Bed_Count", "Quality_Score"]
    for c in numeric_cols:                        # Cell 28 - numeric conversion for EDA
        eda[c] = pd.to_numeric(eda[c], errors="coerce")
    res["numeric_summary"] = eda[numeric_cols].describe().T.reset_index().rename(columns={"index": "Column"})

    res["invalid_beds"] = eda[eda["Bed_Count"] < 0]
    res["invalid_quality"] = eda[(eda["Quality_Score"] < 0) | (eda["Quality_Score"] > 100)]
    res["invalid_coords"] = eda[(eda["Longitude"] < -180) | (eda["Longitude"] > 180) |
                                (eda["Latitude"] < -90) | (eda["Latitude"] > 90)]

    iqr_rows = []
    for c in ["Bed_Count", "Quality_Score", "Latitude", "Longitude"]:
        s = eda[c].dropna()
        if s.empty:
            iqr_rows.append({"Column": c, "Q1": np.nan, "Q3": np.nan, "IQR": np.nan,
                             "Lower_Limit": np.nan, "Upper_Limit": np.nan, "Outlier_Count": 0})
            continue
        q1, q3 = s.quantile(.25), s.quantile(.75)
        iqr = q3 - q1
        lo, hi = q1 - 1.5 * iqr, q3 + 1.5 * iqr
        iqr_rows.append({"Column": c, "Q1": q1, "Q3": q3, "IQR": iqr, "Lower_Limit": lo,
                         "Upper_Limit": hi, "Outlier_Count": int(((s < lo) | (s > hi)).sum())})
    res["iqr"] = pd.DataFrame(iqr_rows)

    issues = []
    for _, r in res["missing"].iterrows():
        if r["Missing_Count"] > 0:
            issues.append({"Issue_Type": "Missing Values", "Column": r["Column"],
                           "Count": int(r["Missing_Count"]), "Percentage": round(float(r["Missing_Percentage"]), 2),
                           "Action_Required": "Handle during preprocessing"})
    n = max(len(eda), 1)
    for label, col, df_, action in [
        ("Invalid Value", "Bed_Count", res["invalid_beds"], "Convert invalid negative values to NA"),
        ("Invalid Value", "Quality_Score", res["invalid_quality"], "Convert values outside 0-100 to NA"),
        ("Invalid Value", "Longitude / Latitude", res["invalid_coords"], "Convert out-of-range coordinates to NA"),
    ]:
        if len(df_) > 0:
            issues.append({"Issue_Type": label, "Column": col, "Count": len(df_),
                           "Percentage": round(len(df_) / n * 100, 2), "Action_Required": action})
    if res["exact_duplicates"] > 0:
        issues.append({"Issue_Type": "Duplicates", "Column": "Entire Record", "Count": res["exact_duplicates"],
                       "Percentage": round(res["exact_duplicates"] / n * 100, 2),
                       "Action_Required": "Remove exact duplicates"})
    res["issue_log"] = pd.DataFrame(issues)
    res["data"] = eda
    return res


def render_eda_results(raw: pd.DataFrame):
    """Single, non-repeating EDA results view (each table / chart appears once)."""
    e = build_eda_results(raw)
    st.caption("Calculated on the combined, schema-standardized data before cleaning, "
               "after treating tokens such as NA, N/A, NULL, None and Unknown as missing.")

    k1, k2, k3, k4 = st.columns(4)
    k1.metric("Rows", f"{e['rows']:,}")
    k2.metric("Columns", e["columns"])
    k3.metric("Exact duplicate rows", e["exact_duplicates"])
    k4.metric("Potential hospital duplicates", f"{e['potential_duplicates']:,}",
              help="Rows repeating the same Hospital_Name + Address + City")

    t1, t2, t3, t4, t5, t6 = st.tabs([
        "🧾 Structure & missing", "🏷️ Categories", "🔢 Numeric & invalid values",
        "🎯 Outliers", "📈 Distributions", "📝 EDA issue log"])

    with t1:
        c1, c2 = st.columns(2)
        with c1:
            st.markdown("**Data types**")
            st.dataframe(e["dtypes"], use_container_width=True, hide_index=True)
        with c2:
            st.markdown("**Unique values per column**")
            st.dataframe(e["unique"], use_container_width=True, hide_index=True)

        st.markdown("**Missing-value analysis (after missing-token standardization)**")
        m1, m2 = st.columns([1, 1.3])
        with m1:
            st.dataframe(e["missing"], use_container_width=True, hide_index=True)
        with m2:
            miss = e["missing"].query("Missing_Count > 0").sort_values("Missing_Count")
            if miss.empty:
                st.success("No missing values were detected in the standardized data.")
            else:
                st.plotly_chart(
                    px.bar(miss, x="Missing_Count", y="Column", orientation="h",
                           hover_data=["Missing_Percentage"], title="Missing values by column"),
                    use_container_width=True)

        st.markdown("**Missing cells by source format**")
        d1, d2 = st.columns([1, 1.3])
        with d1:
            st.dataframe(e["missing_by_source"], use_container_width=True, hide_index=True)
        with d2:
            st.plotly_chart(px.bar(e["missing_by_source"], x="Source_Format", y="Total_Missing_Cells",
                                   text_auto=True, title="Missing cells by source format"),
                            use_container_width=True)

    with t2:
        st.write("Frequency of each value, including missing values, to reveal inconsistent labels "
                 "(for example `in county`, `OUT COUNTY`, `General`, `Emergency`).")
        cats = list(e["categorical"].items())
        for i in range(0, len(cats), 2):
            cols = st.columns(2)
            for col, (name, tbl) in zip(cols, cats[i:i + 2]):
                with col:
                    st.markdown(f"**{name}**")
                    st.dataframe(tbl, use_container_width=True, hide_index=True)
        g1, g2 = st.columns(2)
        with g1:
            st.plotly_chart(px.bar(e["categorical"]["County_Status"], x="County_Status", y="Count",
                                   text_auto=True, title="County status distribution"),
                            use_container_width=True)
        with g2:
            st.plotly_chart(px.bar(e["categorical"]["Hospital_Type"], x="Hospital_Type", y="Count",
                                   text_auto=True, title="Hospital type distribution"),
                            use_container_width=True)

    with t3:
        st.markdown("**Descriptive statistics (numeric fields)**")
        st.dataframe(e["numeric_summary"], use_container_width=True, hide_index=True)
        for title, key, msg in [
            ("Invalid bed count (negative)", "invalid_beds", "negative bed-count record(s)"),
            ("Invalid quality score (outside 0–100)", "invalid_quality", "invalid quality-score record(s)"),
            ("Invalid coordinates", "invalid_coords", "invalid coordinate record(s)"),
        ]:
            st.markdown(f"**{title}**")
            tbl = e[key]
            if tbl.empty:
                st.success(f"0 {msg} found.")
            else:
                st.warning(f"{len(tbl)} {msg} found.")
                st.dataframe(tbl, use_container_width=True, hide_index=True)

    with t4:
        st.markdown("**IQR outlier detection**")
        st.dataframe(e["iqr"], use_container_width=True, hide_index=True)
        st.markdown("**Box plots**")
        data = e["data"]
        cols_ = ["Bed_Count", "Quality_Score", "Latitude", "Longitude"]
        for i in range(0, len(cols_), 2):
            cc_ = st.columns(2)
            for col, name in zip(cc_, cols_[i:i + 2]):
                with col:
                    s = data[name].dropna()
                    if s.empty:
                        st.info(f"No valid data to plot for {name}.")
                    else:
                        st.plotly_chart(px.box(s.to_frame(name), y=name, title=f"Outlier analysis - {name}"),
                                        use_container_width=True)

    with t5:
        data = e["data"]
        h1, h2 = st.columns(2)
        with h1:
            s = data["Bed_Count"].dropna()
            if not s.empty:
                st.plotly_chart(px.histogram(s.to_frame("Bed_Count"), x="Bed_Count", nbins=30,
                                             title="Distribution of hospital bed count"),
                                use_container_width=True)
        with h2:
            s = data["Quality_Score"].dropna()
            if not s.empty:
                st.plotly_chart(px.histogram(s.to_frame("Quality_Score"), x="Quality_Score", nbins=30,
                                             title="Distribution of quality score"),
                                use_container_width=True)
        src = e["categorical"]["Source_Format"]
        st.plotly_chart(px.pie(src, names="Source_Format", values="Count", hole=.45,
                               title="Records by source format"), use_container_width=True)

    with t6:
        st.write("Every problem found during EDA and the action that preprocessing will take.")
        if e["issue_log"].empty:
            st.success("No data-quality issues were logged.")
        else:
            st.dataframe(e["issue_log"], use_container_width=True, hide_index=True)


def page_eda():
    st.header("🔎 Exploratory Data Analysis (EDA)")
    tab_method, tab_results = st.tabs(["📖 EDA steps", "📊 Results (after implementation)"])

    with tab_method:
        st.write("EDA is performed on the **combined, schema-standardized data before cleaning**, "
                 "so that real data-quality problems are visible.")
        st.dataframe(pd.DataFrame([
            {"EDA step": "Structure review", "What it does": "Data type and unique-value count of every column", "Output": "Tables"},
            {"EDA step": "Missing-value analysis", "What it does": "Counts and percentage of missing cells per column and per source format", "Output": "Tables + bar charts"},
            {"EDA step": "Category distribution", "What it does": "Frequency of County_Status, Hospital_Type, Record_Status and Source_Format", "Output": "Tables + bar charts"},
            {"EDA step": "Descriptive statistics", "What it does": "Count, mean, std, min, quartiles, max for numeric fields", "Output": "Summary table"},
            {"EDA step": "Invalid-value check", "What it does": "Negative beds, quality scores outside 0–100, out-of-range coordinates", "Output": "Tables"},
            {"EDA step": "Outlier detection (IQR)", "What it does": "Flags values outside Q1 − 1.5·IQR and Q3 + 1.5·IQR", "Output": "Outlier table + box plots"},
            {"EDA step": "Distributions", "What it does": "Histograms of Bed_Count and Quality_Score; records by source format", "Output": "Histograms + donut chart"},
            {"EDA step": "Issue log", "What it does": "Every problem found and the preprocessing action planned", "Output": "Table"},
        ]), use_container_width=True, hide_index=True)
        st.subheader("About the IQR outlier rule")
        st.markdown("A value is treated as an outlier if it is below **Q1 − 1.5 × IQR** or above **Q3 + 1.5 × IQR**, "
                    "where IQR = Q3 − Q1. Outliers are reported for review; they are not deleted.")

    with tab_results:
        result = require_result()
        if not result:
            return
        render_eda_results(result["raw"])


# ---------------------------------------------------------------------
# 6. IMPLEMENTATION
# ---------------------------------------------------------------------
def page_implementation():
    st.header("▶️ Implementation")
    st.write("Run the full pipeline: parse → map to common schema → combine → preprocess → validate → export. "
             "Files are processed in memory and are not saved to disk.")

    source = st.radio("Data source", ["Use the project dataset files", "Upload my own files"], horizontal=True)

    files = []
    if source == "Use the project dataset files":
        files = sample_files()
        if files:
            st.success(f"{len(files)} dataset file(s) ready: " + ", ".join(f.name for f in files))
        else:
            st.error("No dataset files found. Make sure the `data/` folder with the six files is in your repository.")
    else:
        files = st.file_uploader(
            "Choose one or more files",
            type=["csv", "json", "xml", "html", "htm", "kml", "kmz"],
            accept_multiple_files=True,
            help="Supported formats: CSV, JSON, XML, HTML, KML, KMZ.",
        ) or []

    if st.button("▶️ Run pipeline", type="primary", disabled=not files):
        if process_files(files):
            st.success("Pipeline completed. Results are shown below, and in the **Preprocessing** and **EDA** pages.")

    result = st.session_state.get("hospital_result")
    if not result:
        return

    st.divider()
    t_over, t_data, t_val, t_exp = st.tabs(["📊 Overview", "🗂️ Integrated data", "✅ Validation", "⬇️ Export"])

    with t_over:
        raw = result["raw"]
        k1, k2, k3, k4, k5 = st.columns(5)
        k1.metric("Files processed", len(result["parsed"]))
        k2.metric("Combined records", f"{len(raw):,}")
        k3.metric("Schema fields", len(COMMON_COLUMNS))
        k4.metric("Missing cells before", f"{result['before_after'].loc[result['before_after']['Metric']=='Missing cells','Before'].iloc[0]:,}")
        k5.metric("Validation checks passed", f"{int((result['validation']['Status']=='PASS').sum())}/{len(result['validation'])}")
        left, right = st.columns([1.05, 1])
        with left:
            st.subheader("Records by file")
            st.dataframe(result["source_summary"], use_container_width=True, hide_index=True)
            fig = px.bar(result["source_summary"], x="File", y="Input_Rows", color="Source_Format",
                         title="Input records per file", text_auto=True)
            st.plotly_chart(fig, use_container_width=True)
        with right:
            st.subheader("Before vs after")
            st.dataframe(result["before_after"], use_container_width=True, hide_index=True)
        st.caption("Counts are calculated from the files processed in this session.")

    with t_data:
        raw = result["raw"]
        st.caption("Combined data after schema mapping and before preprocessing.")
        search = st.text_input("Filter rows by text", key="raw_filter")
        view = raw
        if search:
            mask = raw.astype("string").apply(lambda col: col.str.contains(search, case=False, na=False)).any(axis=1)
            view = raw.loc[mask]
        st.dataframe(view, use_container_width=True, hide_index=True, height=460)
        st.download_button("Download standardized CSV", raw.to_csv(index=False).encode("utf-8-sig"),
                           "hospital_standardized_raw.csv", "text/csv")

    with t_val:
        st.dataframe(result["validation"], use_container_width=True, hide_index=True)
        failures = result["validation"].query("Status == 'FAIL'")
        if failures.empty:
            st.success("All configured validation checks passed for this dataset.")
        else:
            st.warning(f"{len(failures)} validation check(s) still report invalid values. Review them before using the dataset.")
        st.caption("A PASS only means the specific check passed. Missing descriptive fields can remain in the final data by design.")

    with t_exp:
        st.write("Download a formatted Excel workbook with the standardized data, cleaned data, schema mapping, "
                 "missing-value report, outlier analysis, before/after comparison, preprocessing log, "
                 "validation report and imputation report.")
        excel_bytes = create_excel_report(
            result["raw"], result["clean"], result["source_summary"], result["mapping"], result["missing"],
            result["validation"], result["imputation"], result["actions"], result["before_after"],
            result["dtypes"], result["outliers"],
        )
        st.download_button(
            "📥 Download full Excel report", data=excel_bytes,
            file_name="Hospital_Data_Final_EDA_Preprocessing.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary",
        )
        st.download_button("Download schema mapping CSV", result["mapping"].to_csv(index=False).encode("utf-8-sig"),
                           "hospital_schema_mapping.csv", "text/csv")
        st.download_button("Download validation CSV", result["validation"].to_csv(index=False).encode("utf-8-sig"),
                           "hospital_validation_report.csv", "text/csv")


# ---------------------------------------------------------------------
# Sidebar navigation
# ---------------------------------------------------------------------
PAGES = {
    "ℹ️  About Project": page_about_project,
    "🏗️  Architecture": page_architecture,
    "📚  About Dataset": page_about_dataset,
    "🔎  EDA": page_eda,
    "🧹  Preprocessing": page_preprocessing,
    "▶️  Implementation": page_implementation,
}

with st.sidebar:
    st.markdown("""
<div class="side-brand">
  <div class="logo">🏥</div>
  <div>
    <div class="title">Intelligent Data<br>Contract Framework</div>
    <div class="sub">INTEGRATE · CLEAN · EXPORT</div>
  </div>
</div>
<div class="side-label">Navigation</div>
""", unsafe_allow_html=True)

    choice = st.radio("Navigate", list(PAGES.keys()), label_visibility="collapsed")

    st.divider()
    if st.session_state.get("hospital_result"):
        n_files = len(st.session_state["hospital_result"]["parsed"])
        st.markdown(f'<div class="status-pill"><span class="dot on"></span>'
                    f'<span>Pipeline ready · {n_files} file(s) processed</span></div>',
                    unsafe_allow_html=True)
    else:
        st.markdown('<div class="status-pill"><span class="dot off"></span>'
                    '<span>Pipeline not run yet</span></div>', unsafe_allow_html=True)
    st.markdown('<div class="side-foot">Tip: run <b>Implementation</b> first to see live results in the '
                'Preprocessing and EDA pages.</div>', unsafe_allow_html=True)

PAGES[choice]()
